#!/usr/bin/env python3
"""
Contratos - TX · gera data/contratos.json a partir da planilha
"CONTROLE DE VIGÊNCIA DE CONTRATOS GERAL".

Uso:  python3 scripts/build_data.py caminho/planilha.xlsx data/contratos.json

O que faz:
  - lê a aba de contratos vigentes e as abas de rescindidos/não renovados;
  - padroniza valores (R$ 1.234,56 / 1234.56), datas e nomes de secretaria;
  - identifica a unidade pelo prefixo do número (2, 3, 4, 5);
  - remove linhas repetidas (mesmo número, credor e valor na mesma aba);
  - aceita colunas opcionais, se existirem no cabeçalho:
        CNPJ / CPF do credor  -> campo "doc"
        LINK / MINUTA / NEXTCLOUD -> campo "minuta"
"""
import sys, json, re, unicodedata, datetime as dt
import openpyxl

UNIDADES = {
    "2": "Prefeitura",
    "3": "Saúde",
    "4": "Educação",
    "5": "Promoção Social",
}

# nome canônico da secretaria  <-  variações encontradas na planilha
SECRETARIAS = {
    "ADMINISTRAÇÃO": ["ADM", "ADMINISTRACAO", "ADMINISTRAÇÃO"],
    "SAÚDE": ["SAUDE", "SAÚDE"],
    "EDUCAÇÃO": ["EDUCACAO", "EDUCAÇÃO", "EDUCAÇAO"],
    "INFRAESTRUTURA": ["INFRA", "INFRAESTRUTURA"],
    "PROMOÇÃO SOCIAL": ["PROMOCAO SOCIAL", "PROMOÇÃO SOCIAL", "PROMOÇÃO", "PROMOCAO",
                        "ASSISTÊNCIA SOCIAL", "ASSISTENCIA SOCIAL"],
    "CULTURA E TURISMO": ["CULTURA E TURISMO", "CULTURA", "TURISMO"],
    "ESPORTE E LAZER": ["ESPORTE E LAZER", "ESPORTE"],
    "GABINETE DO PREFEITO": ["GABINETE DO PREFEITO", "GABINETE"],
    "SEGURANÇA E TRANSPORTE": ["SEGURANÇA E TRANSPORTE", "SEGURANCA E TRANSPORTE",
                               "SEGURANÇA", "SEGURANCA"],
    "PROCURADORIA (PGM)": ["PGM", "PROCURADORIA"],
    "AGRICULTURA": ["AGRICULTURA", "AGRICULTURA, PECUÁRIA E ABASTECIMENTO",
                    "AGRICULTURA, PECUARIA E ABASTECIMENTO"],
}


SEC_POR_UNIDADE = {"3": "SAÚDE", "4": "EDUCAÇÃO", "5": "PROMOÇÃO SOCIAL"}


def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


_SEC_MAP = {}
for canon, variants in SECRETARIAS.items():
    for v in variants:
        _SEC_MAP[sem_acento(v).upper()] = canon


def norm_sec(v):
    if not v:
        return ""
    s = re.sub(r"\s+", " ", str(v)).strip().upper()
    return _SEC_MAP.get(sem_acento(s), s)


def clean(v):
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s.replace("\\&", "&")


def money(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).replace("R$", "").replace("\xa0", "").replace(" ", "").strip()
    if not s or s in {"-", "—"}:
        return None
    if "," in s and "." in s:
        # o separador decimal é o que aparece por último
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None


def date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def parse_numero(num):
    """'2-1023-2025' -> ('2', 1023, 2025)."""
    m = re.match(r"^\s*(\d)\s*-\s*0*(\d+)\s*-\s*(\d{4})\s*$", num or "")
    if not m:
        return None, None, None
    return m.group(1), int(m.group(2)), int(m.group(3))


def find_header(ws):
    for r in range(1, 8):
        vals = [clean(c.value).upper() for c in ws[r]]
        if any("CONTRATO" in v for v in vals) and any("CREDOR" in v for v in vals):
            return r, vals
    return None, None


def col(vals, *keys, exclude=()):
    for i, v in enumerate(vals):
        vv = sem_acento(v)
        if any(sem_acento(k) in vv for k in keys) and not any(sem_acento(e) in vv for e in exclude):
            return i
    return None


def read_sheet(ws, situacao_padrao):
    hr, h = find_header(ws)
    if not hr:
        return [], 0
    c = {
        "num": col(h, "N° DO CONTRATO", "Nº DO CONTRATO", "CONTRATO"),
        "credor": col(h, "CREDOR"),
        "valor": col(h, "VALOR"),
        "mod": col(h, "MODALIDADE"),
        "ini": col(h, "VIGÊNCIA INICIAL", "DATA DE ASSINATURA"),
        "fim": col(h, "VIGÊNCIA FINAL", "DATA DE VIGÊNCIA"),
        "adt": col(h, "ADITIVO"),
        "sec": col(h, "SECRETARIA"),
        "cid": col(h, "CIDADE"),
        "obj": col(h, "OBJETO"),
        "sit": col(h, "SITUÇÃO", "SITUAÇÃO", "SITUACAO"),
        "doc": col(h, "CNPJ", "CPF"),
        "minuta": col(h, "MINUTA", "LINK", "NEXTCLOUD"),
    }
    # aba de vigentes não tem coluna de secretaria com esse nome nas rescindidas (fica na coluna I sem título)
    if c["sec"] is None and len(h) > 8:
        c["sec"] = 8

    def g(row, k):
        i = c.get(k)
        return row[i] if i is not None and i < len(row) else None

    out, seen, dups = [], set(), 0
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        raw_num = g(row, "num")
        num = clean(int(raw_num) if isinstance(raw_num, float) and raw_num.is_integer() else raw_num)
        if not num or sem_acento(num).upper().startswith("NAO"):
            continue
        credor = clean(g(row, "credor"))
        valor = money(g(row, "valor"))
        key = (num, credor.upper(), valor)
        if key in seen:
            dups += 1
            continue
        seen.add(key)
        un, seq, ano = parse_numero(num)
        ini, fim, adt = date(g(row, "ini")), date(g(row, "fim")), date(g(row, "adt"))
        fim_ef = max([d for d in (fim, adt) if d], default=None)
        rec = {
            "n": num,
            "un": un or "",
            "seq": seq,
            "ano": ano,
            "cr": credor,
            "doc": re.sub(r"\D", "", clean(g(row, "doc"))),
            "v": valor,
            "mod": clean(g(row, "mod")).replace("—", "-"),
            "ini": ini,
            "fim": fim,
            "adt": adt,
            "fe": fim_ef,
            "sec": norm_sec(g(row, "sec")),
            "cid": clean(g(row, "cid")).upper(),
            "obj": clean(g(row, "obj")),
            "sit": clean(g(row, "sit")).upper() or situacao_padrao,
            "minuta": clean(g(row, "minuta")),
        }
        if rec["mod"] in {"-", ""}:
            rec["mod"] = ""
        # secretaria em branco nas abas antigas: deduz pela unidade dos fundos
        if not rec["sec"] and un in SEC_POR_UNIDADE:
            rec["sec"] = SEC_POR_UNIDADE[un]
        out.append(rec)
    return out, dups


def main(xlsx, dest):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    vigentes, encerrados, avisos = [], [], []
    for ws in wb.worksheets:
        nome = ws.title.strip().upper()
        if "RESCIND" in nome:
            recs, d = read_sheet(ws, "ENCERRADO")
            for r in recs:
                r["aba"] = ws.title.strip()
            encerrados += recs
        else:
            recs, d = read_sheet(ws, "VIGENTE")
            vigentes += recs
        if d:
            avisos.append(f"{d} linha(s) repetida(s) ignorada(s) na aba '{ws.title.strip()}'")
    sem_mod = sum(1 for r in vigentes if not r["mod"])
    if sem_mod:
        avisos.append(f"{sem_mod} contrato(s) vigente(s) sem nº de modalidade")
    sem_un = [r["n"] for r in vigentes if not r["un"]]
    if sem_un:
        avisos.append("Número fora do padrão U-NNN-AAAA: " + ", ".join(sem_un))
    data = {
        "gerado_em": dt.datetime.now().strftime("%Y-%m-%dT%H:%M"),
        "fonte": "CONTROLE DE VIGÊNCIA DE CONTRATOS GERAL",
        "unidades": UNIDADES,
        "avisos": avisos,
        "vigentes": vigentes,
        "encerrados": encerrados,
    }
    with open(dest, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"vigentes={len(vigentes)} encerrados={len(encerrados)}")
    for a in avisos:
        print("aviso:", a)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
